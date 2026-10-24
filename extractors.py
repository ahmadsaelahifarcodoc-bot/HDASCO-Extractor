"""HDASCO extraction helpers. Site-specific callbacks can change, so scripts validate the page before acting."""
import json
import re
import time
import xml.etree.ElementTree as ET
from pathlib import Path
from urllib.parse import urlparse, parse_qs
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

BR_TABLE = "ctl00_ContentPlaceHolder1_dgBRList"
BL_TABLE = "ctl00_ContentPlaceHolder1_dgBLList"
CONTAINER_TABLE = "ctl00_ContentPlaceHolder1_dgMoreBLList"


def _page_table(cdp, table_id):
    return cdp.evaluate(f"""(() => {{
      const t=document.getElementById({json.dumps(table_id)});
      if(!t) return null;
      return {{headers:[...t.querySelectorAll('tr:first-child th, tr:first-child td')].map(x=>x.innerText.trim()),
        rows:[...t.querySelectorAll('tr')].slice(1).map(tr=>[...tr.querySelectorAll('th,td')].map(td=>td.innerText.trim()))}};
    }})()""")


def _save_table(table, output, sheet_name, log):
    if not table or not table.get("rows"):
        raise RuntimeError("No data rows found. Check that you are logged in and the correct HDASCO list page is open.")
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    headers = table.get("headers") or [f"Column {i+1}" for i in range(max(map(len, table["rows"])))]
    ws.append(headers)
    for row in table["rows"]:
        ws.append(row)
    _format_sheet(ws)
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
    log(f"Saved {len(table['rows'])} rows to {output}")
    return f"{len(table['rows'])} rows; {output}"


def _format_sheet(ws):
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="234A72")
    for col in ws.columns:
        vals = [str(c.value or "") for c in list(col)[:200]]
        width = min(42, max(12, max((len(v) for v in vals), default=10) + 2))
        ws.column_dimensions[get_column_letter(col[0].column)].width = width


def _click_pagination(cdp, page_number):
    # Tries common ASP.NET/Anthem page-link patterns. A page is clicked only if a matching link exists.
    return cdp.evaluate(f"""(() => {{
      const n={int(page_number)};
      const links=[...document.querySelectorAll('a')];
      const a=links.find(x=>x.innerText.trim()===String(n) &&
        (x.closest('table')?.id==={json.dumps(BL_TABLE)} ||
         x.getAttribute('href')?.includes('Page$'+n) ||
         x.getAttribute('href')?.includes('Page%24'+n) ||
         x.getAttribute('onclick')?.includes('Page$'+n)));
      if(!a) return false;
      a.click(); return true;
    }})()""")


def extract_br_list(cdp, pages, delay, export_dir, log):
    if "BRList.aspx" not in cdp.current_url():
        cdp.navigate("https://my.hdasco.com/Packages/Documentation/BRList.aspx")
        time.sleep(2)
    data = []
    headers = None
    for page in range(1, pages + 1):
        table = _page_table(cdp, BR_TABLE)
        if table and table.get("rows"):
            headers = headers or table.get("headers")
            data.extend(table["rows"])
            log(f"BR List: collected page {page}, {len(table['rows'])} displayed rows.")
        else:
            raise RuntimeError("BR List table not found. Sign in to HDASCO and open BR List.")
        if page < pages:
            if not _click_pagination(cdp, page + 1):
                log(f"Could not find a reliable page {page+1} link. Stopping after current page.")
                break
            time.sleep(max(delay, 1.0))
    return _save_table({"headers": headers, "rows": data}, export_dir / "BR_List.xlsx", "BR List", log)


def download_bl_xmls(cdp, pages, delay, export_dir, log):
    if "BLList.aspx" not in cdp.current_url():
        cdp.navigate("https://my.hdasco.com/Packages/Documentation/BLList.aspx")
        time.sleep(2)
    xml_dir = export_dir / "XML"
    xml_dir.mkdir(parents=True, exist_ok=True)
    total = 0
    seen = set()
    for page in range(1, pages + 1):
        buttons = cdp.evaluate(f"""(() => {{
          const t=document.getElementById({json.dumps(BL_TABLE)});
          if(!t) return [];
          return [...t.querySelectorAll('input[type=submit],input[type=button],button')]
            .filter(b=>(b.value||b.innerText||'').trim().toUpperCase()==='GET XML')
            .map(b=>({{id:b.id,name:b.name,value:b.value}}));
        }})()""") or []
        if not buttons:
            raise RuntimeError("No GET XML buttons found. Confirm you are logged in and the BL List table is visible.")
        for ix, btn in enumerate(buttons, 1):
            key = btn.get("id") or btn.get("name") or f"page{page}-row{ix}"
            if key in seen:
                continue
            seen.add(key)
            cdp.evaluate(f"""(() => {{
              const b=document.getElementById({json.dumps(btn.get('id',''))}) ||
                [...document.querySelectorAll('input,button')].find(x=>x.name==={json.dumps(btn.get('name',''))});
              if(!b) return false; b.click(); return true;
            }})()""")
            time.sleep(max(delay, 1.0))
            xml_cdp = None
            try:
                # HDASCO commonly opens ViewBLXMLFile.aspx in a NEW Chrome tab.
                xml_cdp = cdp.connect_matching_url("ViewBLXMLFile.aspx", timeout_seconds=max(10, int(delay) + 8))
                current = xml_cdp.current_url()
                xml_text = xml_cdp.evaluate("new XMLSerializer().serializeToString(document)") or ""
                if not xml_text.strip().startswith("<"):
                    xml_text = xml_cdp.evaluate("document.documentElement.outerHTML || document.documentElement.innerText || ''") or ""
                match = re.search(r"<Data(?:\s|>)", xml_text)
                if match:
                    xml_text = xml_text[match.start():]
                if not xml_text.strip().startswith("<"):
                    raise RuntimeError("The XML tab opened, but its document content could not be read.")
                query_data = parse_qs(urlparse(current).query).get("data", [""])[0]
                # Prefer the BL number in the XML itself; fall back to the query parameter or row number.
                bl_match = re.search(r"<BL(?:No|Number)?>([^<]+)</BL(?:No|Number)?", xml_text, re.I)
                candidate = bl_match.group(1).strip() if bl_match else (query_data or f"BL_page{page}_row{ix}")
                name = re.sub(r"[^A-Za-z0-9._-]+", "_", candidate).strip("._") or f"BL_page{page}_row{ix}"
                if not name.lower().endswith(".xml"):
                    name += ".xml"
                path = xml_dir / name
                if path.exists():
                    path = xml_dir / f"{path.stem}_{total+1}{path.suffix}"
                path.write_text(xml_text, encoding="utf-8")
                # Verify saved content is well-formed before reporting success.
                ET.parse(path)
                total += 1
                log(f"Saved XML {total}: {path.name}")
            except Exception as exc:
                log(f"Could not save XML for GET XML row {ix}: {exc}")
            finally:
                if xml_cdp:
                    xml_cdp.close()
            time.sleep(0.3)
        if page < pages:
            if not _click_pagination(cdp, page + 1):
                log(f"Could not find page {page+1}; stopping after current page.")
                break
            time.sleep(max(delay, 1.0))
    if total == 0:
        raise RuntimeError("No XML files were saved. Check that the XML tabs opened and that you are logged in to HDASCO.")
    return f"Saved {total} XML file(s) in {xml_dir}"


def _text(node, path):
    found = node.find(path)
    return (found.text or "").strip() if found is not None else ""


def _all_text(node):
    return {child.tag: (child.text or "").strip() for child in list(node)}


def export_xml_folder(xml_dir, export_dir, log):
    xml_files = sorted(Path(xml_dir).glob("*.xml"))
    if not xml_files:
        raise RuntimeError(f"No XML files found in {xml_dir}. Download BL XML files first.")
    summary, containers, packages, pricing = [], [], [], []
    for file in xml_files:
        try:
            root = ET.parse(file).getroot()
        except Exception as e:
            log(f"Skipped unreadable XML {file.name}: {e}")
            continue
        manifest = root.find(".//Manifest")
        if manifest is None:
            manifest = root
        base = _all_text(manifest)
        bl_records = manifest.findall(".//BL")
        if not bl_records:
            # Some manifests have a BLS wrapper and BL children.
            bl_records = manifest.findall(".//BLS/BL")
        if not bl_records:
            bl_records = [manifest]
        for bl in bl_records:
            b = _all_text(bl)
            common = {**base, **b}
            bl_no = common.get("BLNo") or common.get("BLNumber") or common.get("BL") or file.stem
            summary.append({"SourceFile": file.name, **common, "BL_Number": bl_no})
            for container in bl.findall(".//Container"):
                c = _all_text(container)
                cno = c.get("ContainerNo") or c.get("ContainerNumber") or c.get("No") or ""
                containers.append({"SourceFile": file.name, "BL_Number": bl_no, **c})
                pkg_nodes = container.findall(".//Package") + container.findall(".//Packages/Package")
                if not pkg_nodes:
                    packages.append({"SourceFile": file.name, "BL_Number": bl_no, "ContainerNo": cno,
                                     "Package": "", "Note": "No Package element detected"})
                for pkg in pkg_nodes:
                    packages.append({"SourceFile": file.name, "BL_Number": bl_no, "ContainerNo": cno, **_all_text(pkg)})
            for node in bl.findall(".//Pricing/*"):
                pricing.append({"SourceFile": file.name, "BL_Number": bl_no, "PricingType": node.tag, **_all_text(node)})
            # Flexible fallback for common pricing/cost structures.
            for node in bl.findall(".//Price") + bl.findall(".//Rate") + bl.findall(".//Charge"):
                pricing.append({"SourceFile": file.name, "BL_Number": bl_no, "PricingType": node.tag, **_all_text(node)})
    output = Path(export_dir) / "HDASCO_BL_Export.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "BL Summary"
    _write_dicts(ws, summary)
    for name, rows in [("Container Details", containers), ("Cargo Packages", packages), ("Pricing Details", pricing)]:
        _write_dicts(wb.create_sheet(name), rows)
    wb.save(output)
    log(f"Excel workbook saved: {output}")
    return f"{len(xml_files)} XML file(s) processed; workbook: {output}"


def _write_dicts(ws, records):
    keys = []
    for record in records:
        for key in record:
            if key not in keys:
                keys.append(key)
    if not keys:
        keys = ["No records detected"]
        records = [{keys[0]: ""}]
    ws.append(keys)
    for record in records:
        ws.append([record.get(k, "") for k in keys])
    _format_sheet(ws)


def extract_container_history(cdp, pages, delay, export_dir, log):
    if "BLList.aspx" not in cdp.current_url():
        cdp.navigate("https://my.hdasco.com/Packages/Documentation/BLList.aspx")
        time.sleep(2)
    # Site-specific detail controls require live-session validation. Export the visible expanded table
    # when it is already present, rather than guessing a callback and silently corrupting data.
    table = _page_table(cdp, CONTAINER_TABLE)
    if not table or not table.get("rows"):
        raise RuntimeError(
            "Container history table is not currently visible. Open a BL's Details first on the HDASCO page, "
            "then run this function. Automated opening of each Details callback must be verified against your live session."
        )
    return _save_table(table, export_dir / "Container_History.xlsx", "Container History", log)
